"""Pro batch processing API.

POST /api/batches/                         create a batch from several files (Pro, server-checked)
GET  /api/batches/                         the signed-in user's recent batches
GET  /api/batches/<id>                     live status/progress (computed from the real job rows)
GET  /api/batches/<id>/files/<job>         one file's status + cleaning summary/audit
GET  /api/batches/<id>/files/<job>/download   one cleaned file (repeatable until retention ends)
GET  /api/batches/<id>/download            ZIP of COMPLETED files + batch_report.csv
POST /api/batches/<id>/cancel              cancel files that have not started
POST /api/batches/<id>/retry               re-run FAILED files (Pro)

Plan checks happen here, from the subscription record. Nothing in the request (fields, headers,
JS) can grant Pro. Every file passes the same checks as a single upload.
"""
import io
import json
import logging
import os
import shutil
import tempfile
import uuid

from flask import Blueprint, jsonify, redirect, request, send_file

import db as metadata_db
import jobs.service as job_service
import storage as object_storage
from accounts import auth
from accounts.entitlements import max_upload_bytes, pro_required
from accounts.store import get_store
from cleaning.options import OptionsError, parse_options
from cleaning.rules import DEFAULT_RULES, RULE_DISPATCH
from config import Config
from jobs import batches
from utils.file_handler import (
    allowed_file, is_valid_job_id, safe_display_name, set_job_owner, validate_upload_content,
)
from utils.observability import get_request_id
from utils.session import current_session_hash

batches_bp = Blueprint("batches", __name__)
logger = logging.getLogger("omixa.batches")
MB = 1024 * 1024


def limits() -> dict:
    per_file = max_upload_bytes(True)
    if Config.BATCH_MAX_FILE_MB:
        per_file = min(per_file, Config.BATCH_MAX_FILE_MB * MB)
    cells = min(Config.MAX_CELLS, Config.BATCH_MAX_CELLS_PER_FILE) if Config.BATCH_MAX_CELLS_PER_FILE else Config.MAX_CELLS
    return {
        "max_files": Config.BATCH_MAX_FILES,
        "max_total_mb": Config.BATCH_MAX_TOTAL_MB,
        "max_file_mb": per_file // MB,
        "max_cells_per_file": cells,
        "max_total_cells": Config.BATCH_MAX_TOTAL_CELLS,
        "concurrent_files": Config.BATCH_MAX_CONCURRENT_FILES,
        "retention_hours": round(Config.BATCH_RETENTION_SECONDS / 3600, 1),
        "formats": sorted(Config.ALLOWED_EXTENSIONS),
    }


@batches_bp.get("/limits")
def get_limits():
    return jsonify(limits())


# --------------------------------------------------------------------------- helpers

def estimate_cells(path: str, ext: str):
    """Cheap pre-scan (no full parse). Returns None when it cannot be estimated (.xls);
    the pipeline still enforces the per-file cap while processing."""
    try:
        if ext == "csv":
            import csv
            with open(path, "rb") as fh:
                lines, last = 0, b"\n"
                for chunk in iter(lambda: fh.read(1024 * 1024), b""):
                    lines += chunk.count(b"\n")
                    last = chunk[-1:]
            if last != b"\n":
                lines += 1
            with open(path, "r", encoding="utf-8", errors="replace", newline="") as fh:
                header = next(csv.reader(fh), [])
            return max(lines, 1) * max(len(header), 1)
        if ext == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True)
            try:
                ws = wb.worksheets[0]
                return int((ws.max_row or 1) * (ws.max_column or 1))
            finally:
                wb.close()
    except Exception:
        return None
    return None


def _owned_batch(batch_id: str):
    """(batch, error_response). Only the owning signed-in user may see a batch; others get 404."""
    user, err = auth.current_user_record()
    if err:
        return None, err
    if not is_valid_job_id(batch_id):
        return None, (jsonify({"error": "Batch not found"}), 404)
    batch = metadata_db.get_batch(batch_id)
    if not batch or batch["owner_uid"] != auth.current_uid():
        return None, (jsonify({"error": "Batch not found"}), 404)
    return batch, None


def _file_job(batch, job_id):
    if not is_valid_job_id(job_id):
        return None
    job = metadata_db.get_job(job_id)
    if not job or job.get("batch_id") != batch["batch_id"]:
        return None
    return job


def _check_choices(opts: dict, where: str = ""):
    """Validates one set of cleaning choices (rules, resolutions, leave_as_is, options). Returns (clean, error)."""
    out = {}
    rules = opts.get("rules")
    if rules is not None:
        if not isinstance(rules, list) or not all(isinstance(r, str) for r in rules):
            return None, f"{where}'rules' must be a list of rule names"
        unknown = [r for r in rules if r not in RULE_DISPATCH]
        if unknown:
            return None, f"{where}Unknown rule(s): {', '.join(unknown)}"
        out["rules"] = rules
    res = opts.get("resolutions")
    if res is not None:
        if not (isinstance(res, list) and len(res) <= 500 and all(isinstance(r, dict) for r in res)):
            return None, f"{where}'resolutions' must be a list of objects"
        out["resolutions"] = res
    las = opts.get("leave_as_is")
    if las is not None:
        if not (isinstance(las, list) and len(las) <= 1000 and all(isinstance(c, str) for c in las)):
            return None, f"{where}'leave_as_is' must be a list of column name strings"
        out["leave_as_is"] = las
    if opts.get("options") is not None:
        try:
            out["options"] = parse_options(opts["options"])
        except OptionsError as exc:
            return None, f"{where}{exc}"
    return out, None


def _validate_opts(opts):
    """Batch-wide choices, same vocabulary as /api/process. Returns (opts, error)."""
    if not isinstance(opts, dict):
        return None, "'options' must be an object"
    clean, err = _check_choices(opts)
    if err:
        return None, err
    if opts.get("standardize_column_names") is True and clean.get("rules") is None:
        # Opt-in header renaming: the default rule set plus column_names (which must run first).
        clean["rules"] = ["column_names"] + [r for r in DEFAULT_RULES if r != "column_names"]
    if not isinstance(opts.get("has_header", True), bool):
        return None, "'has_header' must be true or false"
    clean["has_header"] = opts.get("has_header", True)
    if opts.get("profile_id") is not None and not isinstance(opts["profile_id"], str):
        return None, "'profile_id' must be text"
    clean["profile_id"] = opts.get("profile_id")
    return clean, None


def _parse_options():
    raw = request.form.get("options")
    if not raw:
        return {"has_header": True}, None
    try:
        opts = json.loads(raw)
    except ValueError:
        return None, "'options' must be valid JSON"
    return _validate_opts(opts)


def _build_spec(uid, opts, max_cells):
    """(spec, error_response). The spec is what every file's worker runs; Pro is decided server-side."""
    profile = None
    if opts.get("profile_id"):
        store = get_store()
        profile = store.get_profile(uid, opts["profile_id"]) if store is not None else None
        if not profile:
            return None, (jsonify({"error": "Quality Profile not found"}), 404)
    return {"kind": "clean", "rules": opts.get("rules"), "resolutions": opts.get("resolutions"),
            "has_header": opts.get("has_header", True), "pro": True, "uid": uid, "profile": profile,
            "cleaning_profile": None, "leave_as_is": opts.get("leave_as_is"),
            "options": opts.get("options") or {}, "max_cells": max_cells}, None


# --------------------------------------------------------------------------- create

@batches_bp.post("/")
@pro_required
def create_batch():
    if Config.PROCESSING_MODE != "queue":
        return jsonify({"error": "Batch processing needs Omixa's processing workers, which are not enabled on this server."}), 503

    uid = auth.current_uid()
    lim = limits()
    files = request.files.getlist("files")
    if not files:
        return jsonify({"error": "Choose at least one file."}), 400
    if len(files) > lim["max_files"]:
        return jsonify({"error": f"A batch can contain at most {lim['max_files']} files.", "limits": lim}), 400

    opts, err = _parse_options()
    if err:
        return jsonify({"error": err}), 400
    spec, bad = _build_spec(uid, opts, lim["max_cells_per_file"])
    if bad:
        return bad
    draft = request.form.get("start", "true").strip().lower() == "false"

    if metadata_db.count_active_batches(uid) >= Config.BATCH_MAX_ACTIVE_PER_USER:
        resp = jsonify({"error": f"You already have {Config.BATCH_MAX_ACTIVE_PER_USER} unfinished batch(es). "
                                 "Wait for one to finish (or cancel it) before starting another."})
        resp.headers["Retry-After"] = "30"
        return resp, 429
    session_hash = current_session_hash()
    try:
        job_service.admit(uid, session_hash, True)
    except job_service.AdmissionError as exc:
        resp = jsonify({"error": exc.message, "status": "rejected"})
        if exc.retry_after:
            resp.headers["Retry-After"] = str(exc.retry_after)
        return resp, exc.status

    tmpdir = tempfile.mkdtemp(prefix="omx-batch-")
    accepted, skipped = [], []
    total_bytes = total_cells = 0
    try:
        for f in files:
            raw_name = f.filename or ""
            shown = safe_display_name(raw_name, raw_name.rsplit(".", 1)[-1].lower()) if "." in raw_name else (raw_name[:80] or "(no name)")

            def skip(reason, shown=shown):
                skipped.append({"filename": shown, "reason": reason})

            if not raw_name:
                skip("File has no name.")
                continue
            if not allowed_file(raw_name):
                skip("Unsupported file type. Use CSV, XLSX or XLS.")
                continue
            ext = raw_name.rsplit(".", 1)[1].lower()
            bad = validate_upload_content(f, ext)       # magic bytes, macros, zip-bomb guard
            if bad:
                skip(bad)
                continue
            tmp = os.path.join(tmpdir, f"{uuid.uuid4().hex}.{ext}")  # generated name, never the client's
            f.save(tmp)
            size = os.path.getsize(tmp)
            if size == 0:
                skip("File is empty.")
                continue
            if size > lim["max_file_mb"] * MB:
                skip(f"File is larger than the {lim['max_file_mb']} MB per-file limit.")
                continue
            total_bytes += size
            if total_bytes > lim["max_total_mb"] * MB:
                return jsonify({"error": f"The batch is larger than {lim['max_total_mb']} MB in total.", "limits": lim}), 413
            cells = estimate_cells(tmp, ext)
            if cells is not None and cells > lim["max_cells_per_file"]:
                total_bytes -= size
                skip(f"File has more than {lim['max_cells_per_file']:,} cells.")
                continue
            total_cells += cells or 0
            if total_cells > lim["max_total_cells"]:
                return jsonify({"error": f"The batch has more than {lim['max_total_cells']:,} cells in total.", "limits": lim}), 413
            accepted.append({"tmp": tmp, "ext": ext, "name": safe_display_name(raw_name, ext), "size": size})

        if not accepted:
            return jsonify({"error": "None of the files could be accepted.", "skipped": skipped}), 400

        batch_id = str(uuid.uuid4())
        if not metadata_db.create_batch(batch_id, uid, session_hash, {"spec": spec, "draft": draft}, skipped, get_request_id()):
            return jsonify({"error": "Omixa is temporarily unavailable. Please try again."}), 503

        storage = object_storage.get_storage()
        created = []
        try:
            for seq, a in enumerate(accepted):
                job_id = str(uuid.uuid4())              # internal id: the only thing used for storage
                storage.put_path(object_storage.source_key(job_id, a["ext"]), a["tmp"])
                created.append(job_id)
                if not object_storage.is_remote():
                    set_job_owner(job_id, session_hash)
                if not metadata_db.record_upload(job_id, session_hash, a["ext"], a["name"], a["size"],
                                                 owner_uid=uid, request_id=get_request_id()):
                    raise RuntimeError("record_upload failed")
                metadata_db.attach_job_to_batch(job_id, batch_id, seq)
        except Exception:
            logger.exception("batch %s: could not store files", batch_id)
            for jid in created:
                storage.delete_job(jid)
            metadata_db.cancel_batch_pending(batch_id)
            metadata_db.finalize_batch(batch_id, "failed")
            return jsonify({"error": "Could not save the uploaded files. Please try again."}), 503
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)

    metadata_db.set_batch_totals(batch_id, len(accepted), total_bytes)
    metadata_db.audit(uid, "batch.created", batch_id, request_id=get_request_id(),
                      meta={"files": len(accepted), "skipped": len(skipped), "total_bytes": total_bytes})
    batches.advance(batch_id)  # release the first files to the queue; the request returns right away
    batch = metadata_db.get_batch(batch_id)
    return jsonify(batches.batch_view(batch, metadata_db.batch_jobs(batch_id))), 202


# --------------------------------------------------------------------------- read

@batches_bp.get("/")
def list_batches():
    user, err = auth.current_user_record()
    if err:
        return err
    out = []
    for b in metadata_db.list_batches_for(auth.current_uid()):
        v = batches.batch_view(b, metadata_db.batch_jobs(b["batch_id"]))
        v.pop("files", None)
        out.append(v)
    return jsonify({"batches": out})


@batches_bp.get("/<batch_id>")
def get_batch(batch_id):
    batch, err = _owned_batch(batch_id)
    if err:
        return err
    jobs = metadata_db.batch_jobs(batch_id)
    return jsonify(batches.batch_view(batch, jobs))


@batches_bp.get("/<batch_id>/files/<job_id>")
def get_file(batch_id, job_id):
    batch, err = _owned_batch(batch_id)
    if err:
        return err
    job = _file_job(batch, job_id)
    if not job:
        return jsonify({"error": "File not found"}), 404
    out = job_service.public_status(job)
    out["filename"] = job.get("original_filename")
    out["status"] = batches.file_status(job)
    return jsonify(out)


# --------------------------------------------------------------------------- download

@batches_bp.get("/<batch_id>/files/<job_id>/download")
def download_file(batch_id, job_id):
    batch, err = _owned_batch(batch_id)
    if err:
        return err
    job = _file_job(batch, job_id)
    if not job or batches.file_status(job) != "completed":
        return jsonify({"error": "This file is not available for download."}), 404
    ext = job.get("cleaned_ext") or "csv"
    name = batches.cleaned_download_name(job)
    storage = object_storage.get_storage()
    key = object_storage.cleaned_key(job_id, ext)
    if not storage.exists(key):
        return jsonify({"error": "This file is no longer available."}), 404
    url = storage.presigned_get(key, name)
    if url:
        return redirect(url, code=302)
    tmpdir = tempfile.mkdtemp(prefix="omx-dl-")
    try:
        path = os.path.join(tmpdir, f"out.{ext}")
        storage.get_to_path(key, path)
        with open(path, "rb") as fh:
            data = io.BytesIO(fh.read())
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)
    mimetype = {"csv": "text/csv", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}.get(ext, "application/octet-stream")
    return send_file(data, mimetype=mimetype, as_attachment=True, download_name=name)


@batches_bp.get("/<batch_id>/download")
def download_all(batch_id):
    batch, err = _owned_batch(batch_id)
    if err:
        return err
    jobs = metadata_db.batch_jobs(batch_id)
    if not any(batches.file_status(j) == "completed" for j in jobs):
        return jsonify({"error": "No completed files to download yet."}), 409
    spool = batches.build_zip(batch, jobs)
    return send_file(spool, mimetype="application/zip", as_attachment=True,
                     download_name=f"omixa-batch-{batch_id[:8]}.zip")


# --------------------------------------------------------------------------- control

@batches_bp.post("/<batch_id>/cancel")
def cancel(batch_id):
    batch, err = _owned_batch(batch_id)
    if err:
        return err
    n = metadata_db.cancel_batch_pending(batch_id)
    metadata_db.audit(auth.current_uid(), "batch.cancelled", batch_id, meta={"cancelled_files": n})
    batches.advance(batch_id)
    return jsonify(batches.batch_view(metadata_db.get_batch(batch_id), metadata_db.batch_jobs(batch_id)))


@batches_bp.post("/<batch_id>/retry")
@pro_required
def retry(batch_id):
    batch, err = _owned_batch(batch_id)
    if err:
        return err
    n = metadata_db.reset_batch_failed(batch_id)
    if n:
        metadata_db.audit(auth.current_uid(), "batch.retried", batch_id, meta={"files": n})
        batches.advance(batch_id)
    return jsonify(batches.batch_view(metadata_db.get_batch(batch_id), metadata_db.batch_jobs(batch_id)))


# --------------------------------------------------------------------------- draft flow: analyse -> choose -> start

def _draft_batch(batch_id):
    batch, err = _owned_batch(batch_id)
    if err:
        return None, None, err
    opts = json.loads(batch.get("options_json") or "{}")
    if batch["finalized"] or not opts.get("draft"):
        return None, None, (jsonify({"error": "This batch has already started."}), 409)
    return batch, opts, None


@batches_bp.post("/<batch_id>/analyze")
@pro_required
def analyze(batch_id):
    """Checks every file (quality score, findings, flagged values) without changing anything."""
    batch, opts, err = _draft_batch(batch_id)
    if err:
        return err
    body = request.get_json(silent=True) or {}
    hh = body.get("has_header", (opts.get("spec") or {}).get("has_header", True))
    if not isinstance(hh, bool):
        return jsonify({"error": "'has_header' must be true or false"}), 400
    opts.setdefault("spec", {})["has_header"] = hh
    opts["analyzing"] = True
    metadata_db.set_batch_options(batch_id, opts)
    metadata_db.reset_batch_failed(batch_id)
    batches.advance(batch_id)
    return jsonify(batches.batch_view(metadata_db.get_batch(batch_id), metadata_db.batch_jobs(batch_id)))


@batches_bp.get("/<batch_id>/files/<job_id>/analysis")
def file_analysis(batch_id, job_id):
    batch, err = _owned_batch(batch_id)
    if err:
        return err
    job = _file_job(batch, job_id)
    a = batches.analysis_of(job) if job else None
    if not a:
        return jsonify({"error": "This file has not been analysed yet."}), 404
    return jsonify({"job_id": job_id, "filename": job.get("original_filename"), "has_header": a.get("has_header"),
                    "report": a.get("report"), "recommendations": a.get("recommendations"),
                    "validation": a.get("validation"), "flagged_values": a.get("flagged_values")})


@batches_bp.post("/<batch_id>/start")
@pro_required
def start(batch_id):
    """Starts a draft batch with the user's choices: batch-wide, plus optional per-file overrides
    {per_file: {job_id: {rules, resolutions, leave_as_is, options, skip}}}."""
    batch, old, err = _draft_batch(batch_id)
    if err:
        return err
    body = request.get_json(silent=True)
    if not isinstance(body, dict):
        return jsonify({"error": "Send the cleaning choices as a JSON object."}), 400
    uid = auth.current_uid()
    opts, e = _validate_opts(body)
    if e:
        return jsonify({"error": e}), 400
    spec, bad = _build_spec(uid, opts, limits()["max_cells_per_file"])
    if bad:
        return bad
    jobs = {j["job_id"]: j for j in metadata_db.batch_jobs(batch_id)}
    per_file, skip = {}, []
    raw_pf = body.get("per_file") or {}
    if not isinstance(raw_pf, dict) or len(raw_pf) > Config.BATCH_MAX_FILES:
        return jsonify({"error": "'per_file' must be an object keyed by file id"}), 400
    for jid, ov in raw_pf.items():
        if jid not in jobs:
            return jsonify({"error": "Unknown file in 'per_file'"}), 400
        if not isinstance(ov, dict):
            return jsonify({"error": "Each 'per_file' entry must be an object"}), 400
        if ov.get("skip") is True:
            skip.append(jid)
            continue
        clean, e = _check_choices(ov, f"{jobs[jid].get('original_filename')}: ")
        if e:
            return jsonify({"error": e}), 400
        if clean:
            per_file[jid] = clean
    metadata_db.reset_batch_failed(batch_id)         # files whose analysis failed are still cleaned
    for jid in skip:
        metadata_db.cancel_job_in_batch(batch_id, jid)
    metadata_db.set_batch_options(batch_id, {"spec": spec, "per_file": per_file, "draft": False})
    metadata_db.audit(uid, "batch.started", batch_id, request_id=get_request_id(),
                      meta={"files": len(jobs) - len(skip), "skipped_by_user": len(skip),
                            "rules": len(spec.get("rules") or DEFAULT_RULES)})
    batches.advance(batch_id)
    return jsonify(batches.batch_view(metadata_db.get_batch(batch_id), metadata_db.batch_jobs(batch_id))), 202
