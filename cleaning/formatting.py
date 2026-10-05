"""
Per-column style profiling for Excel exports (see export/exporter.py).

Context (16-17.09.2026 testing report + survey): the MVP forced every
cell in the output to one blanket font/size (e.g. everything became
Calibri 11pt regardless of the source), stripped bold/italic
wholesale, and reset every column to one default width, even when a
column's own formatting was clearly intentional (e.g. a lexicography
file's IPA transcription column consistently in a different font).
The fix isn't to preserve every individual cell's exact original
style (that would just re-import whatever inconsistency the user
was trying to clean up in the first place) - it's to standardize
each column to *its own* majority style rather than a single
sheet-wide default, and to size each column to fit its own content
rather than forcing one uniform width.

This module only ever reads. It never touches the DataFrame or the
source file; export/exporter.py is what applies the profile it
returns when it rebuilds the cleaned sheet.

CSV sources carry no cell-level formatting at all (there's no such
thing as a "bold" CSV cell), so this module is simply never invoked
for a CSV upload - export_dataframe skips straight to the plain
df.to_csv path.
"""

from __future__ import annotations
from collections import Counter
from typing import Any, Dict, List, Optional

DEFAULT_FONT_NAME = "Calibri"
DEFAULT_FONT_SIZE = 11.0
MIN_COL_WIDTH = 8.0
MAX_COL_WIDTH = 60.0
WIDTH_PADDING = 2.0
# Excel's own hard ceiling on a column's width, used only as a
# sanity bound on an EXPLICIT user-set width (a corrupted or
# out-of-range stored value shouldn't crash the export). This is
# deliberately much looser than MIN/MAX_COL_WIDTH above, which only
# apply to widths OMIXA itself computes by autofitting content -
# manual intent (e.g. a column purposely made much narrower than its
# content would need) is respected as-is, not squeezed into that
# aesthetic range.
_EXCEL_MAX_WIDTH = 255.0

# Column-width policy (resolves the three open questions the
# 16.09.2026 report left for its own "Solution 2": what if one
# column is purposely much shorter than similar-looking neighbours,
# what if most columns differ only slightly, what if the longest
# value is very long). The report's own "Solution 1" already answers
# all three at once without needing to compare columns against each
# other at all: if the user (or whatever app last saved the file)
# EXPLICITLY set a column's width, that's manual intent and is
# respected outright, no similarity heuristics needed - a "purposely
# much shorter" column simply keeps being shorter. Only a column
# nobody ever touched (Excel/Sheets' own untouched default) falls
# back to autofit-by-content, still clamped to MIN/MAX so a single
# very long value can't blow a column out to an unusable width.


def extract_column_style_profile(
    ws, n_cols: int, has_header: bool
) -> List[Optional[Dict[str, Any]]]:
    """
    ws: the ORIGINAL (uncleaned) openpyxl worksheet, read before
    export/exporter.py deletes and rebuilds it.
    n_cols: how many source columns to profile, matching the source
    sheet's own column positions (0-indexed) - see
    processing/pipeline.py's `source_positions`, which is what lets
    the exporter later say "source column 3" no matter how many
    columns a resolution dropped, or how column_names renamed the
    survivors, between read and export.
    has_header: when True, row 1 is excluded from the majority count
    (a header's own styling, e.g. bold, shouldn't be counted as
    "the normal style for this column's data"), matching the same
    has_header the caller already gave read_source() at read time.

    Returns a list of length n_cols. Each entry is either None (the
    column had no non-empty data cells to profile, e.g. it's fully
    blank - callers should leave that column's formatting alone) or:
        {"font_name": str, "size": float, "bold": bool,
         "italic": bool, "width": float}
    where font_name/size/bold/italic are whichever combination is
    most common among that column's own data cells, and width is the
    source column's own explicitly-set width if it had one, else
    autofit to that column's own longest value (clamped to a sane
    min/max) - see the module-level comment above on why this needs
    no cross-column comparison.
    """
    from openpyxl.utils import get_column_letter

    start_row = 2 if has_header else 1
    counters = [Counter() for _ in range(n_cols)]
    max_lens = [0] * n_cols
    header_lens = [0] * n_cols

    if has_header:
        header_row = next(ws.iter_rows(min_row=1, max_row=1), ())
        for cell in header_row:
            idx = cell.column - 1
            if 0 <= idx < n_cols and cell.value is not None:
                header_lens[idx] = len(str(cell.value))

    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            idx = cell.column - 1
            if idx < 0 or idx >= n_cols or cell.value is None:
                continue
            font = cell.font
            key = (
                font.name or DEFAULT_FONT_NAME,
                float(font.size) if font.size else DEFAULT_FONT_SIZE,
                bool(font.bold),
                bool(font.italic),
            )
            counters[idx][key] += 1
            length = len(str(cell.value))
            if length > max_lens[idx]:
                max_lens[idx] = length

    profile: List[Optional[Dict[str, Any]]] = []
    for i in range(n_cols):
        if not counters[i]:
            profile.append(None)
            continue
        (font_name, size, bold, italic), _ = counters[i].most_common(1)[0]

        dim = ws.column_dimensions.get(get_column_letter(i + 1))
        if dim is not None and dim.customWidth and dim.width:
            # Explicit width the user (or their spreadsheet app) set
            # on purpose - respected as-is (including "purposely much
            # narrower/wider than the content"), only guarded against
            # a corrupted/out-of-range stored value.
            width = min(_EXCEL_MAX_WIDTH, max(0.5, float(dim.width)))
        else:
            content_width = max(max_lens[i], header_lens[i]) + WIDTH_PADDING
            width = min(MAX_COL_WIDTH, max(MIN_COL_WIDTH, content_width))

        profile.append(
            {
                "font_name": font_name,
                "size": size,
                "bold": bold,
                "italic": italic,
                "width": width,
            }
        )
    return profile
