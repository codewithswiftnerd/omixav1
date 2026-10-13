"""Data-integrity regression tests: deliberately messy files go through the REAL API (upload -> process ->
download) and every assertion is about the data that comes back and what was reported about it.

Principles checked: valid values are not touched, safe formatting is normalised per cell, ambiguous or
unreadable values are preserved and flagged (never guessed, never zeroed), no row disappears silently, the
reports are recomputed from the final dataset, and the exported file really contains the cleaned data.
"""
import io

import openpyxl
import pandas as pd
import pytest

from cleaning.options import OptionsError, parse_options
from cleaning.quality_report import generate_report
from export.verify import verify_export


# --------------------------------------------------------------------------------------------- helpers
def _upload(client, payload: bytes, name: str) -> str:
    r = client.post("/api/upload/", data={"file": (io.BytesIO(payload), name)}, content_type="multipart/form-data")
    assert r.status_code in (200, 201), r.get_json()
    return r.get_json()["job_id"]


def run_csv(client, text: str, *, options=None, rules=None):
    """Returns (summary, exported CSV as an all-text DataFrame, raw exported text)."""
    job = _upload(client, text.encode("utf-8"), "messy.csv")
    body = {}
    if options is not None:
        body["options"] = options
    if rules is not None:
        body["rules"] = rules
    r = client.post(f"/api/process/{job}", json=body)
    assert r.status_code == 200, r.get_json()
    summary = r.get_json()["summary"]
    dl = client.get(f"/api/download/{job}")
    assert dl.status_code == 200
    raw = dl.data.decode("utf-8-sig")
    return summary, pd.read_csv(io.StringIO(raw), dtype=str, keep_default_na=False), raw


def col_report(summary, name):
    return next(c for c in summary["integrity_report"]["columns"] if c["column"] == name)


def changes(summary, column=None):
    cells = summary["change_log"]["cell_changes"]
    return [c for c in cells if column is None or c["column"] == column]


MESSY = (
    "Customer ID,Name,Email,Age,Quantity,Price,Score,City,Status,Signup Date,Phone\n"
    "C001,John Doe,john@x.com,25,2,1500,45,Lagos,Active,2024-01-15,08031234567\n"
    'C002,mary jane,mary@x.com, 32 ,5,"₦2,500",67,lagos,active,15/02/2024,08031234568\n'
    "C003,Peter Obi,peter@x,ERROR,UNKNOWN,invalid,N/A,LAGOS,ACTIVE,2024-13-45,0803\n"
    "C004,Ada Eze,ada@x.com,thirty,8,3500,82,Lagos ,Inactive,2024-03-01,+2348031234569\n"
    "C005,Ada Eze,ada2@x.com,41,-3,3500,90,Lagoss,Active,2024-03-02,08031234570\n"
    "C006,John Doe,john@x.com,25,2,1500,45,Lagos,Active,2024-01-15,08031234567\n"
    "C007,Sam,sam@x.com,-5,1,-200,70,Abuja,Active,2090-01-01,08031234571\n"
    "C008,Tim,tim@x.com,250,3,100,NULL,Abuja,Active,2024-05-05,08031234572\n"
    'C009,Ken,ken@x.com,comma,4,"1,234",55,Abuja,Active,2024-05-06,08031234573\n'
    "C010,Joy,joy@x.com,forty-five,6,500,60,Abuja,Active,2024-05-07,08031234574\n"
)


# ------------------------------------------------------------------ 1. valid values stay unchanged
def test_valid_numbers_are_not_rewritten_as_floats_when_a_marker_sits_in_the_column(client):
    # pandas used to read "N/A" as NaN, turn the column into float64 and export 45 as "45.0"
    s, out, raw = run_csv(client, "id,score\n1,45\n2,N/A\n3,82\n4,67\n")
    assert out["score"].tolist() == ["45", "", "82", "67"]
    assert "45.0" not in raw


def test_untouched_text_survives_exactly_leading_zeros_trailing_zeros_and_whole_floats(client):
    s, out, _ = run_csv(client, "id,phone,price\n1,08031234567,45.50\n2,08031234568,7\n3,08031234569,3.25\n")
    assert out["phone"].tolist() == ["08031234567", "08031234568", "08031234569"]
    assert out["price"].tolist() == ["45.50", "7", "3.25"]      # never "45.5" / "7.0" for cells nobody changed
    assert s["export_verification"]["verified"] is True


# ------------------------------------------------------------------ 2. missing values
def test_every_missing_marker_and_blank_becomes_missing_and_the_original_is_in_the_audit_trail(client):
    csv = "id,note,qty\n1,a,5\n2, ,N/A\n3,b,NULL\n4,   ,NA\n5,c,None\n6,d,7\n"
    s, out, _ = run_csv(client, csv)
    assert out["qty"].tolist() == ["5", "", "", "", "", "7"]
    assert out["note"].tolist() == ["a", "", "b", "", "c", "d"]
    originals = {c["original"] for c in changes(s, "qty")}
    assert {"N/A", "NULL", "NA", "None"} <= originals                       # nothing swallowed before the audit
    assert s["integrity_report"]["missing_values"]["converted_from_placeholders"] >= 4
    assert s["integrity_report"]["missing_values"]["filled"] == 0           # never invents data


def test_error_and_unknown_are_not_treated_identically(client):
    # ERROR in a number column is NOT a recognised marker by default: preserved and flagged.
    # UNKNOWN in a quantity column is blanked, but UNKNOWN in a categorical Status column is a real label.
    csv = ("id,age,qty,status\n1,25,2,Active\n2,32,5,Unknown\n3,ERROR,UNKNOWN,Active\n4,41,8,Unknown\n"
           "5,30,1,Active\n6,28,3,Active\n")
    s, out, _ = run_csv(client, csv)
    assert out["age"].tolist()[2] == "ERROR"
    assert out["qty"].tolist()[2] == ""
    assert out["status"].tolist().count("Unknown") == 2
    assert col_report(s, "age")["needs_review_after"] >= 1


def test_configured_markers_turn_error_into_missing_only_in_a_clearly_numeric_column(client):
    csv = "id,age,remark\n1,25,fine\n2,32,fine\n3,ERROR,error in scan\n4,41,ERROR\n5,30,fine\n6,28,fine\n"
    s, out, _ = run_csv(client, csv, options={"missing_markers": ["error"]})
    assert out["age"].tolist()[2] == ""                                     # converted, audited
    assert any(c["original"] == "ERROR" and c["new"] is None for c in changes(s, "age"))
    assert out["remark"].tolist()[3] == "ERROR"                             # a text column keeps its words
    assert out["remark"].tolist()[2] == "error in scan"


# ------------------------------------------------------------------ 3/4. numeric normalisation, mixed types
def test_one_bad_value_no_longer_blocks_the_rest_of_the_column(client):
    s, out, _ = run_csv(client, "id,price\n1,1500\n2,\"₦2,500\"\n3,invalid\n4,3500\n5,\" 12 \"\n")
    assert out["price"].tolist() == ["1500", "2500", "invalid", "3500", "12"]
    assert "0" not in out["price"].tolist()                                 # arbitrary text is never turned into zero
    assert col_report(s, "price")["needs_review_after"] == 1
    ch = [c for c in changes(s, "price") if c["original"] == "₦2,500"]
    assert ch and ch[0]["new"] == 2500 and ch[0]["approval"] == "automatic" and ch[0]["rule"] == "numeric_text_cleaning"


def test_words_in_a_number_column_are_flagged_and_converted_only_when_the_rule_is_enabled(client):
    csv = "id,age\n1,25\n2,30\n3,comma\n4,40\n5,forty-five\n6,50\n"
    s, out, _ = run_csv(client, csv)                                        # default: nothing guessed
    assert out["age"].tolist() == ["25", "30", "comma", "40", "forty-five", "50"]
    assert col_report(s, "age")["needs_review_after"] == 2

    s2, out2, _ = run_csv(client, csv, options={"number_words": True})
    assert out2["age"].tolist() == ["25", "30", "comma", "40", "45", "50"]  # unambiguous word read, "comma" never guessed
    ch = [c for c in changes(s2, "age") if c["original"] == "forty-five"][0]
    assert ch["new"] == 45 and ch["approval"] == "user_selected"            # the user switched the rule on: not "automatic"
    assert col_report(s2, "age")["needs_review_after"] == 1                 # only "comma" remains
    assert s2["integrity_report"]["status"] != "successfully_cleaned"


@pytest.mark.parametrize("body,expected_cell,flagged", [
    # decimal-comma evidence elsewhere in the column: "1,234" could be 1.234 or 1234 -> not guessed
    ("id,amount\n1,\"3,45\"\n2,\"1,234\"\n3,\"12,5\"\n4,7\n", "1,234", 1),
    # thousands evidence ("12,500.50") -> unambiguous
    ("id,amount\n1,\"12,500.50\"\n2,\"1,234\"\n3,800\n", "1234", 0),
    # a recognised currency symbol makes the thousands reading unambiguous
    ("id,amount\n1,\"₦2,500\"\n2,\"1,234\"\n3,800\n", "1234", 0),
])
def test_ambiguous_thousands_versus_decimal_comma_is_decided_by_the_column_not_guessed(client, body, expected_cell, flagged):
    s, out, _ = run_csv(client, body)
    assert out["amount"].tolist()[1] == expected_cell
    assert col_report(s, "amount")["needs_review_after"] == flagged
    if flagged:
        assert s["integrity_report"]["fully_clean"] is False


def test_explicit_number_format_decides_the_ambiguous_value(client):
    body = "id,amount\n1,\"3,45\"\n2,\"1,234\"\n3,\"12,5\"\n4,7\n"
    s, out, _ = run_csv(client, body, options={"number_format": "decimal_comma"})
    assert out["amount"].tolist() == ["3.45", "1.234", "12.5", "7"]
    assert col_report(s, "amount")["needs_review_after"] == 0


# ------------------------------------------------------------------ 5. impossible / suspicious values
def test_impossible_and_suspicious_values_are_flagged_as_warnings_and_never_changed(client):
    s, out, _ = run_csv(client, MESSY)
    ages = out["Age"].tolist()
    assert "-5" in ages and "250" in ages                                   # preserved exactly
    assert out["Quantity"].tolist()[4] == "-3"
    sus = {(f["column"], f["issue"]) for f in s["quality_report"]["findings"]}
    assert ("Age", "suspicious_values") in sus
    assert ("Quantity", "suspicious_values") in sus
    assert ("Signup Date", "suspicious_values") in sus or ("Signup Date", "impossible_date") in sus
    age = col_report(s, "Age")
    assert age["counts_after"]["suspicious"] == 2                           # -5 and 250, even though other cells are bad
    lvl = {f["issue"]: f["level"] for f in s["quality_report"]["findings"] if f["column"] == "Age"}
    assert lvl["suspicious_values"] in ("low", "medium")                    # a warning, not a critical error


def test_a_negative_quantity_can_be_declared_legitimate(client):
    csv = "id,quantity\n1,5\n2,-2\n3,8\n4,3\n"
    s, _, _ = run_csv(client, csv)
    assert col_report(s, "quantity")["counts_after"]["suspicious"] == 1
    s2, _, _ = run_csv(client, csv, options={"validation": {"rules": {"quantity": {"allow_negative": True}}}})
    assert col_report(s2, "quantity")["counts_after"]["suspicious"] == 0


def test_invalid_dates_and_malformed_emails_and_phones_are_flagged_not_repaired(client):
    s, out, _ = run_csv(client, MESSY)
    assert out["Signup Date"].tolist()[2] == "2024-13-45"
    assert out["Email"].tolist()[2] == "peter@x"
    assert out["Phone"].tolist()[2] == "0803"
    types = {c["column"]: c["detected_type"] for c in s["integrity_report"]["columns"]}
    assert types["Signup Date"] == "date"                                   # was misread as "phone" (digits and dashes)
    assert types["Phone"] == "phone" and types["Email"] == "email"
    assert col_report(s, "Signup Date")["counts_after"]["invalid"] == 1
    assert col_report(s, "Email")["counts_after"]["invalid"] == 1
    assert col_report(s, "Phone")["counts_after"]["invalid"] == 1


# ------------------------------------------------------------------ 6. categories
def test_case_and_spacing_are_normalised_but_a_possible_misspelling_is_only_suggested(client):
    rows = ["Lagos", "lagos", "LAGOS", "Lagos ", "Lagoss", "Lagos", "Abuja", "Abuja", "Lagos", "Abuja", "Lagos", "Abuja",
            "Lagos", "Abuja", "Lagos", "Lagos"]
    csv = "id,city\n" + "\n".join(f"{i},{c}" for i, c in enumerate(rows, 1)) + "\n"
    s, out, _ = run_csv(client, csv)
    cities = out["city"].tolist()
    assert cities.count("Lagos") == 10 and "Lagoss" in cities               # casing/spacing merged, Lagoss NOT replaced
    assert "lagos" not in cities and "LAGOS" not in cities and "Lagos " not in cities
    f = [x for x in s["quality_report"]["findings"] if x["issue"] == "possible_misspelling" and x["column"] == "city"]
    assert f and f[0]["affected_count"] == 1 and "Lagoss" in f[0]["detail"] + str(f[0].get("sample_values", ["Lagoss"]))
    assert not any(c["original"] == "Lagoss" for c in changes(s, "city"))


def test_a_different_real_word_is_not_suggested_as_a_typo(client):
    csv = "id,status\n" + "\n".join(f"{i},Active" for i in range(1, 8)) + "\n8,Inactive\n"
    s, out, _ = run_csv(client, csv)
    assert "Inactive" in out["status"].tolist()
    assert not any(f["issue"] == "possible_misspelling" for f in s["quality_report"]["findings"])


# ------------------------------------------------------------------ 7. duplicates / silent deletion
def test_same_name_different_email_is_never_a_duplicate_and_exact_copies_are_removed_with_an_audit_entry(client):
    csv = ("name,email\nAda Eze,ada@x.com\nAda Eze,ada2@x.com\nBob,bob@x.com\nBob,bob@x.com\n")
    s, out, _ = run_csv(client, csv)
    assert out["email"].tolist() == ["ada@x.com", "ada2@x.com", "bob@x.com"]    # both Adas kept
    d = s["integrity_report"]["duplicates"]
    assert d["exact_rows_detected"] == 1 and d["exact_rows_removed"] == 1 and d["removal_authorized"] is True
    assert d["potential"]["groups"] == 0 and d["by_key"]["groups"] == 0         # a shared name is not an identity
    removed = s["change_log"]["removed_rows"]
    assert len(removed) == 1 and removed[0]["source_row"] == 5                  # the spreadsheet row that went
    assert s["integrity_report"]["rows"] == {"original": 4, "final": 3, "removed": 1}


def test_key_duplicates_and_shared_identifiers_are_reported_for_review_never_merged(client):
    csv = ("id,name,email\n1,Ada Eze,ada@x.com\n2,Bob,bob@x.com\n3,Ada E.,ada@x.com\n4,Cy,cy@x.com\n")
    s, out, _ = run_csv(client, csv, options={"duplicate_keys": ["email"]})
    assert len(out) == 4                                                        # nothing merged or removed
    d = s["integrity_report"]["duplicates"]
    assert d["by_key"]["groups"] == 1 and d["by_key"]["rows"] == 2 and d["by_key"]["samples"] == [[2, 4]]
    assert d["potential"]["groups"] == 1
    assert s["integrity_report"]["fully_clean"] is False
    assert any("share a key" in w for w in s["integrity_report"]["warnings"])


def test_no_row_is_silently_deleted_and_unticking_duplicates_keeps_every_row(client):
    s, out, _ = run_csv(client, MESSY)
    ir = s["integrity_report"]
    assert len(out) == ir["rows"]["final"]
    assert ir["rows"]["original"] - ir["rows"]["final"] == s["change_log"]["removed_rows_total"] == ir["rows"]["removed"]
    s2, out2, _ = run_csv(client, "name,email\nBob,bob@x.com\nBob,bob@x.com\n", rules=["formatting"])
    assert len(out2) == 2                                                       # removal is opt-in per rule set
    assert s2["integrity_report"]["duplicates"]["removal_authorized"] is False
    assert any("still in the file" in w for w in s2["integrity_report"]["warnings"])


# ------------------------------------------------------------------ 8/9. honest reporting
def test_the_report_is_recomputed_from_the_final_data_and_never_claims_a_dirty_file_is_clean(client):
    s, out, raw = run_csv(client, MESSY)
    ir = s["integrity_report"]
    assert ir["status"] in ("partially_cleaned", "requires_manual_review") and ir["fully_clean"] is False
    assert ir["invalid_values"]["requiring_manual_review"] > 0 and ir["unresolved_issues"]
    assert ir["rows"]["original"] == 10 and ir["columns_count"] == {"original": 11, "final": 11}
    assert col_report(s, "Customer ID")["status"] == "successfully_cleaned"
    assert col_report(s, "Age")["status"] in ("partially_cleaned", "requires_manual_review")
    statuses = {c["status"] for c in ir["columns"]}
    assert statuses <= {"successfully_cleaned", "partially_cleaned", "requires_manual_review"}
    assert "NOT fully clean" in s["cleaning_summary"]["text"]

    # independent proof: validating the file the user actually downloads gives exactly the numbers the report claims
    import os, tempfile
    from processing.pipeline import read_source
    from cleaning.validation import validate_dataframe
    with tempfile.NamedTemporaryFile("w", suffix=".csv", delete=False, encoding="utf-8") as f:
        f.write(raw)
    try:
        totals = validate_dataframe(read_source(f.name, has_header=True), None, include_examples=False)["totals"]
    finally:
        os.unlink(f.name)
    after = ir["invalid_values"]["by_kind_after"]
    assert {k: totals[k] for k in ("invalid", "suspicious", "unresolved")} == after


def test_a_genuinely_clean_file_is_reported_as_fully_clean(client):
    s, out, _ = run_csv(client, "id,name,age\n1,Ada,25\n2,Bob,32\n3,Cy,41\n")
    ir = s["integrity_report"]
    assert ir["status"] == "successfully_cleaned" and ir["fully_clean"] is True
    assert out["age"].tolist() == ["25", "32", "41"] and ir["warnings"] == []


def test_quality_score_equals_a_fresh_report_on_the_exported_file(client):
    s, out, raw = run_csv(client, MESSY)
    from processing.pipeline import read_source
    import tempfile, os
    with tempfile.NamedTemporaryFile("w", suffix=".csv", delete=False, encoding="utf-8") as f:
        f.write(raw)
    try:
        fresh = generate_report(read_source(f.name, has_header=True))
    finally:
        os.unlink(f.name)
    assert fresh["score"] == s["quality_report"]["score"]


# ------------------------------------------------------------------ 10. audit trail
def test_every_changed_value_is_traceable_to_its_original(client):
    s, out, _ = run_csv(client, MESSY)
    cells = s["change_log"]["cell_changes"]
    assert cells and s["change_log"]["truncated"] is False
    for c in cells:
        assert {"row", "source_row", "column", "original", "new", "rule", "approval"} <= set(c)
        assert c["approval"] in ("automatic", "user_selected", "user_approved")
    price = [c for c in cells if c["column"] == "Price" and c["original"] == "₦2,500"][0]
    assert price["source_row"] == 3 and price["new"] == 2500                    # row 3 of the uploaded sheet (header = row 1)
    assert any(c["column"] == "Score" and c["original"] == "N/A" and c["new"] is None for c in cells)
    assert any(c["column"] == "Score" and c["original"] == "NULL" for c in cells)


# ------------------------------------------------------------------ 11. export validation
def test_export_is_verified_against_dataset_ledger_and_upload(client):
    s, _, _ = run_csv(client, MESSY)
    ev = s["export_verification"]
    assert ev["verified"] is True
    names = {c["check"]: c["passed"] for c in ev["checks"]}
    assert names == {"readable": True, "header": True, "cells_match": True, "changes_applied": True,
                     "untouched": True, "revalidation": True}


def test_verification_detects_a_file_that_does_not_match_the_cleaned_data(tmp_path):
    df = pd.DataFrame({"a": ["1", "2", "3"], "b": ["x", "y", "z"]})
    bad = tmp_path / "out.csv"
    bad.write_text("a,b\n1,x\n2,TAMPERED\n3,z\n", encoding="utf-8")
    ev = verify_export(out_path=str(bad), ext="csv", cleaned_df=df, kept_index=[0, 1, 2], column_labels={"a": "a", "b": "b"},
                       source_text_df=df.copy(), source_positions=[0, 1], ledger={"cell_changes": [], "truncated": False},
                       validation_after_totals=None)
    assert ev["verified"] is False
    failed = {c["check"] for c in ev["checks"] if c["passed"] is False}
    assert {"cells_match", "untouched"} <= failed
    short = tmp_path / "short.csv"
    short.write_text("a,b\n1,x\n", encoding="utf-8")
    ev2 = verify_export(out_path=str(short), ext="csv", cleaned_df=df, kept_index=[0, 1, 2], column_labels={}, source_text_df=None,
                        source_positions=[0, 1], ledger={}, validation_after_totals=None)
    assert ev2["verified"] is False and ev2["checks"][0]["check"] == "readable"


def test_verification_detects_a_recorded_change_that_never_reached_the_file(tmp_path):
    df = pd.DataFrame({"a": ["99", "2"]})                                       # the dataset (and the ledger) say 99 ...
    f = tmp_path / "o.csv"
    f.write_text("a\n10\n2\n", encoding="utf-8")                               # ... but the file still holds the old 10
    ledger = {"cell_changes": [{"row": 0, "column": "a", "original": "10", "new": 99}], "truncated": False}
    ev = verify_export(out_path=str(f), ext="csv", cleaned_df=df, kept_index=[0, 1], column_labels={"a": "a"},
                       source_text_df=None, source_positions=[0], ledger=ledger, validation_after_totals=None)
    failed = {c["check"] for c in ev["checks"] if c["passed"] is False}
    assert ev["verified"] is False and {"cells_match", "changes_applied"} <= failed


def test_formula_injection_guard_is_reported_as_the_only_dataset_vs_file_difference(client):
    s, out, raw = run_csv(client, "id,note\n1,fine\n2,=SUM(A1:A9)\n3,@cmd\n4,-5\n")
    assert out["note"].tolist()[1] == "'=SUM(A1:A9)" and out["note"].tolist()[3] == "-5"
    ev = s["export_verification"]
    assert ev["verified"] is True and any("apostrophe" in n for n in ev["notes"])


def test_excel_in_excel_out_matches_cleaned_data_and_keeps_numbers_as_numbers(client):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in [["Customer", "Age", "Price", "Phone"], ["Ada", 25, 1500, "08031234567"], ["Bob", " 32 ", "₦2,500", "08031234568"],
                ["Cy", "thirty", 3500.5, "0803"], ["Di", 41, "invalid", "08031234570"], ["Di", 41, "invalid", "08031234570"]]:
        ws.append(row)
    wb.create_sheet("Other").append(["keep", "me"])
    buf = io.BytesIO()
    wb.save(buf)
    job = _upload(client, buf.getvalue(), "messy.xlsx")
    r = client.post(f"/api/process/{job}", json={})
    assert r.status_code == 200
    s = r.get_json()["summary"]
    assert s["export_verification"]["verified"] is True
    out = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/download/{job}").data))
    assert out.sheetnames == ["Data", "Other"]                                  # other sheets carried through
    rows = list(out["Data"].iter_rows(values_only=True))
    assert rows[0] == ("Customer", "Age", "Price", "Phone") and len(rows) == 5  # one exact duplicate removed
    assert rows[2][2] == 2500 and isinstance(rows[2][2], (int, float))          # real number, not "₦2,500"
    assert rows[1][3] == "08031234567" and rows[3][3] == "0803"                 # phones stay text, zeros kept
    assert rows[3][1] == "thirty" and rows[4][2] == "invalid"                   # unreadable values preserved
    assert s["integrity_report"]["rows"]["removed"] == 1


# ------------------------------------------------------------------ 12. options + entitlements
@pytest.mark.parametrize("bad", [
    {"number_format": "european"}, {"unknown_option": True}, {"missing_markers": "error"}, {"missing_markers": [""]},
    {"duplicate_keys": []}, {"number_words": "yes"}, {"validation": "strict"},
])
def test_invalid_options_are_rejected_with_a_clear_400(client, bad):
    job = _upload(client, b"a,b\n1,2\n", "x.csv")
    r = client.post(f"/api/process/{job}", json={"options": bad})
    assert r.status_code == 400 and r.get_json()["code"] == "INVALID_OPTIONS"
    with pytest.raises(OptionsError):
        parse_options(bad)


def test_options_never_unlock_pro_rules_for_a_free_user(client):
    job = _upload(client, b"name\nada eze\nBOB\n", "x.csv")
    profile = {"columns": {"name": {"rules": [{"type": "normalize_case", "mode": "title"}]}}}
    r = client.post(f"/api/process/{job}", json={"cleaning_profile": profile, "options": {"number_words": True}})
    assert r.status_code == 402 and r.get_json()["upgrade_required"] is True


def test_omitting_options_behaves_exactly_as_before(client):
    csv = "id,score\n1,45\n2,N/A\n3,82\n"
    a, out_a, _ = run_csv(client, csv)
    b, out_b, _ = run_csv(client, csv, options={})
    assert out_a.equals(out_b) and a["integrity_report"]["status"] == b["integrity_report"]["status"]
