"""
One regression test per bug fixed in the audit. Each docstring says what used to go wrong.
"""
import hashlib
import io
import logging
import os
import re
import zipfile

import pandas as pd
import pytest

import config
from cleaning import detectors
from cleaning.rules import apply_rules
from tests.helpers import norm_nulls

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _upload(client, name, data, ctype="text/csv"):
    return client.post("/api/upload/", data={"file": (io.BytesIO(data), name, ctype)},
                       content_type="multipart/form-data")


def _job_dirs():
    from config import Config
    return [d for d in os.listdir(Config.TEMP_DIR)] if os.path.isdir(Config.TEMP_DIR) else []


# ============================================================ stability
def test_app_module_imports_and_starts_serving():
    """app.py contained two stray-character syntax errors, so the application could not start."""
    import ast
    for rel in ("app.py", "config.py", "db.py"):
        ast.parse(open(os.path.join(ROOT, rel), encoding="utf-8").read())
    import app as app_module
    resp = app_module.app.test_client().get("/api/health")
    assert resp.status_code == 200


def test_every_python_file_in_the_repo_parses():
    import ast
    bad = []
    for base, dirs, files in os.walk(ROOT):
        dirs[:] = [d for d in dirs if d not in {".git", "__pycache__", "node_modules", "_legacy"}]
        for f in files:
            if f.endswith(".py"):
                try:
                    ast.parse(open(os.path.join(base, f), encoding="utf-8").read())
                except SyntaxError as e:
                    bad.append((f, str(e)))
    assert not bad


def test_security_module_shares_the_app_config_class():
    """Tests used importlib.reload(config), which replaced Config for later tests and leaked
    tiny rate limits into the rest of the suite (order-dependent 429s)."""
    import utils.security as sec
    assert sec.Config is config.Config


def test_rate_limit_state_does_not_leak_between_tests(client):
    for _ in range(3):
        assert client.get("/api/health").status_code == 200


def test_fractional_median_does_not_crash_integer_columns():
    """Median-filling a whole-number column whose median is fractional (59.5) raised
    TypeError ("Invalid value '59.5' for dtype 'Int64'") and failed the whole job."""
    df = pd.DataFrame({"satisfaction": ["56%", "44%", "Unknown", "82%", "71%", "63%", "56%"]})
    out, log = apply_rules(df, rules=["missing_token_normalization", "numeric_text_cleaning", "missing_values"])
    assert out["satisfaction"].isna().sum() == 0
    assert 59.5 in out["satisfaction"].tolist() or 56.0 in out["satisfaction"].tolist()
    ints = pd.Series(pd.array([44, 56, None, 82, 71], dtype="Int64"))
    filled, med = detectors.fill_with_median(ints)  # median 63.5: used to raise TypeError
    assert filled.isna().sum() == 0 and med == 63.5


def test_whole_number_median_keeps_the_integer_type():
    s, med = detectors.fill_with_median(pd.Series(pd.array([10, None, 30], dtype="Int64")))
    assert med == 20 and str(s.dtype) == "Int64"


def test_latin1_csv_that_passes_upload_validation_can_be_processed(client):
    """Upload accepts Latin-1 text, but the pipeline read it as UTF-8 and crashed on 'é'."""
    data = "name,city\nJosé,Zürich\nRené,Genève\n".encode("latin-1")
    r = _upload(client, "latin.csv", data)
    assert r.status_code == 201, r.get_json()
    job = r.get_json()["job_id"]
    assert client.get(f"/api/report/{job}").status_code == 200
    p = client.post(f"/api/process/{job}", json={"rules": ["formatting"]})
    assert p.status_code == 200 and p.get_json()["status"] == "completed"
    assert "José" in client.get(f"/api/download/{job}").get_data(as_text=True)


def test_utf8_bom_csv_does_not_pollute_the_first_column_name(client):
    data = "\ufeffname,age\nAnn,30\n".encode("utf-8")
    job = _upload(client, "bom.csv", data).get_json()["job_id"]
    cols = client.get(f"/api/report/{job}").get_json()["report"]["column_types"]
    assert "name" in cols


def test_date_columns_are_recognised_by_word_not_by_substring():
    """'date' in name matched candidate_name / validated_by / mandate / update_note."""
    for not_date in ("candidate_name", "validated_by", "mandate", "update_note"):
        assert not detectors.has_date_name(not_date)
    for is_date in ("Signup Date", "dob", "order_date", "startDate", "birthdate"):
        assert detectors.has_date_name(is_date)


def test_text_column_named_like_a_date_but_with_no_dates_is_not_flagged():
    from cleaning.quality_report import generate_report
    df = pd.DataFrame({"candidate_name": ["Ann", "Bob", "Cy", "Dee"]})
    assert not any(f["issue"] == "ambiguous_date_format" for f in generate_report(df)["findings"])


def test_standardising_dates_never_throws_away_a_time_of_day():
    df = pd.DataFrame({"created_at": ["2024-01-05 10:30:00", "2024-02-06 11:45:00", "2024-03-07 09:15:00"]})
    out, log = apply_rules(df, rules=["date_standardization"])
    assert out["created_at"].tolist() == df["created_at"].tolist()
    assert log["details"]["date_standardization"]["skipped_has_time_component"] == ["created_at"]


def test_two_digit_years_are_not_silently_given_a_century():
    df = pd.DataFrame({"signup_date": ["05/01/24", "06/02/68", "07/03/99"]})
    out, log = apply_rules(df, rules=["date_standardization"])
    assert out["signup_date"].tolist() == df["signup_date"].tolist()
    assert "signup_date" in log["details"]["date_standardization"]["skipped_ambiguous"]


def test_missing_placeholders_are_not_double_counted_as_invalid_values():
    """'Unknown' in gender/email/phone columns was reported both as missing AND as invalid."""
    from cleaning.quality_report import generate_report
    df = pd.DataFrame({"Gender": ["male", "female", "Unknown", "male"],
                       "email": ["a@b.com", "N/A", "c@d.org", "e@f.net"]})
    issues = {f["issue"] for f in generate_report(df)["findings"]}
    assert "unrecognized_gender_value" not in issues and "invalid_email_format" not in issues
    assert "missing_values" in issues


# ============================================================ secrets
def _tracked_text_files():
    skip_dirs = {".git", "__pycache__", "node_modules", "tests", "testing", "_legacy"}
    for base, dirs, files in os.walk(ROOT):
        dirs[:] = [d for d in dirs if d not in skip_dirs]
        for f in files:
            if f == "conftest.py":
                continue
            path = os.path.join(base, f)
            if os.path.getsize(path) > 2_000_000:
                continue
            try:
                yield path, open(path, encoding="utf-8").read()
            except (UnicodeDecodeError, OSError):
                continue


def test_no_credentials_file_is_shipped_in_the_repo():
    """RAILWAY_ENV_VARS.txt contained a real secret key and an admin password in plain text."""
    names = [f for _, _, files in os.walk(ROOT) for f in files]
    assert not [n for n in names if n.upper().startswith("RAILWAY_ENV_VARS")]
    assert not [n for n in names if n in {".env"} or n.endswith(".pem")]


def test_no_secret_looking_values_are_hardcoded_in_repo_files():
    patterns = [
        re.compile(r"OMIXA_SECRET_KEY\s*=\s*[A-Za-z0-9_\-]{24,}"),
        re.compile(r"scrypt:\d+:\d+:\d+\$[A-Za-z0-9+/=]{8,}\$[0-9a-f]{40,}"),
        re.compile(r"pbkdf2:sha\d+:\d+\$"),
        re.compile(r"OMIXA_ADMIN_PASSWORD\s*=\s*\S{6,}"),
    ]
    offenders = []
    for path, text in _tracked_text_files():
        for pat in patterns:
            if pat.search(text):
                offenders.append((os.path.relpath(path, ROOT), pat.pattern[:30]))
    assert not offenders, offenders


def test_gitignore_blocks_credential_dumps():
    gi = open(os.path.join(ROOT, ".gitignore")).read()
    assert "RAILWAY_ENV_VARS" in gi and ".env" in gi


def test_production_refuses_a_secret_key_that_was_previously_committed(monkeypatch):
    compromised = "previously-committed-key-" + "x" * 30
    monkeypatch.setattr(config, "_COMPROMISED_SECRET_SHA256",
                        {hashlib.sha256(compromised.encode()).hexdigest()})
    monkeypatch.setattr(config, "_IS_PRODUCTION", True)
    monkeypatch.setenv("OMIXA_SECRET_KEY", compromised)
    with pytest.raises(RuntimeError, match="compromised"):
        config._secret_key()


def test_production_refuses_a_short_secret_key_and_accepts_a_strong_one(monkeypatch):
    monkeypatch.setattr(config, "_IS_PRODUCTION", True)
    monkeypatch.setenv("OMIXA_SECRET_KEY", "short")
    with pytest.raises(RuntimeError, match="too short"):
        config._secret_key()
    monkeypatch.setenv("OMIXA_SECRET_KEY", "a-long-random-looking-value-0123456789abcdef")
    assert config._secret_key().startswith("a-long")


def test_the_denylist_stores_only_fingerprints_never_the_secret():
    assert config._COMPROMISED_SECRET_SHA256
    assert all(re.fullmatch(r"[0-9a-f]{64}", h) for h in config._COMPROMISED_SECRET_SHA256)


def test_generate_secrets_script_prints_and_never_writes_files(tmp_path, monkeypatch):
    import contextlib
    import importlib.util
    spec = importlib.util.spec_from_file_location("gen", os.path.join(ROOT, "scripts", "generate_secrets.py"))
    gen = importlib.util.module_from_spec(spec); spec.loader.exec_module(gen)
    import tempfile
    empty = tempfile.mkdtemp()  # tmp_path is shared with the autouse db/temp-dir fixtures
    monkeypatch.chdir(empty)
    monkeypatch.setattr("sys.argv", ["generate_secrets.py"])
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        gen.main()
    out = buf.getvalue()
    assert out.startswith("OMIXA_SECRET_KEY=") and len(out.strip().split("=", 1)[1]) >= 48
    assert os.listdir(empty) == []


def test_error_messages_stored_for_admins_are_scrubbed():
    """Exception text can embed server paths, emails and account numbers from the user's data."""
    import db
    msg = db.scrub_error_message("failed at /home/app/tmp/job-1/source.csv for jane@corp.com acct 12345678901")
    assert "/home/app" not in msg and "jane@corp.com" not in msg and "12345678901" not in msg
    assert len(db.scrub_error_message("x" * 5000)) <= 200


# ============================================================ upload / file handling
def test_non_ascii_filename_is_accepted_not_a_500(client):
    """secure_filename('名前.csv') returns 'csv' (no dot), so the old split raised IndexError -> 500
    and left a half-created job folder behind."""
    r = _upload(client, "名前.csv", b"a,b\n1,2\n")
    assert r.status_code == 201, r.get_json()
    assert r.get_json()["filename"].endswith(".csv")


def test_path_traversal_in_filename_is_neutralised(client):
    r = _upload(client, "../../etc/passwd.csv", b"a,b\n1,2\n")
    assert r.status_code == 201
    name = r.get_json()["filename"]
    assert "/" not in name and ".." not in name and name.endswith(".csv")


def test_script_in_filename_is_not_echoed_back_raw(client):
    r = _upload(client, '<script>alert(1)</script>.csv', b"a,b\n1,2\n")
    assert r.status_code == 201
    assert "<" not in r.get_json()["filename"]


def test_failed_save_leaves_no_orphan_job_directory(client, monkeypatch):
    import routes.upload as up
    def boom(*a, **k):
        raise RuntimeError("disk full at /secret/path")
    monkeypatch.setattr(up, "save_upload", boom)
    before = set(_job_dirs())
    r = _upload(client, "x.csv", b"a,b\n1,2\n")
    assert r.status_code == 500
    assert "/secret/path" not in r.get_data(as_text=True)
    assert set(_job_dirs()) == before


def test_delete_job_removes_nested_content_and_ignores_garbage_ids(client):
    from utils.file_handler import create_job_dir, delete_job, job_dir_path
    job = create_job_dir()
    d = job_dir_path(job)
    os.makedirs(os.path.join(d, "nested"))
    open(os.path.join(d, "nested", "leftover.tmp"), "w").write("user data")
    delete_job(job)
    assert not os.path.exists(d)
    for junk in ("../etc", "", "x" * 500, "a/b"):
        delete_job(junk)  # must not raise


def test_xlsx_zip_bomb_is_rejected_before_parsing(client, monkeypatch):
    from config import Config
    monkeypatch.setattr(Config, "MAX_XLSX_UNCOMPRESSED_MB", 1)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("xl/workbook.xml", "<workbook/>")
        z.writestr("xl/worksheets/sheet1.xml", b"0" * (3 * 1024 * 1024))  # 3 MB inflated, a few KB on disk
    r = _upload(client, "bomb.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert r.status_code in (400, 415, 422)
    assert "unreasonable size" in r.get_data(as_text=True)


def test_oversized_row_count_is_rejected_with_a_clear_message(client, monkeypatch):
    from config import Config
    monkeypatch.setattr(Config, "MAX_ROWS", 5)
    rows = "a,b\n" + "\n".join(f"{i},x" for i in range(20))
    job = _upload(client, "big.csv", rows.encode()).get_json()["job_id"]
    r = client.post(f"/api/process/{job}", json={"rules": ["formatting"]})
    assert r.status_code == 422 and "rows" in r.get_json()["error"]
    assert client.get(f"/api/report/{job}").status_code == 422


def test_too_many_columns_is_rejected(client, monkeypatch):
    from config import Config
    monkeypatch.setattr(Config, "MAX_COLUMNS", 3)
    job = _upload(client, "wide.csv", b"a,b,c,d,e\n1,2,3,4,5\n").get_json()["job_id"]
    r = client.get(f"/api/report/{job}")
    assert r.status_code == 422 and "columns" in r.get_json()["error"]


# ============================================================ request handling
def test_spoofed_x_forwarded_for_cannot_bypass_the_rate_limit(client, monkeypatch):
    """The limiter keyed on the client-supplied X-Forwarded-For, so a fresh fake IP per request
    defeated it (and the login brute-force limit)."""
    from config import Config
    import utils.security as sec
    monkeypatch.setattr(Config, "RATE_LIMIT_PER_MINUTE", 3)
    sec._hits.clear()
    codes = [client.get("/api/report/does-not-exist", headers={"X-Forwarded-For": f"10.0.0.{i}"}).status_code
             for i in range(8)]
    assert 429 in codes


def test_cross_origin_post_is_blocked_but_same_origin_and_non_browser_are_allowed(client):
    evil = client.post("/api/upload/", data={"file": (io.BytesIO(b"a\n1\n"), "x.csv")},
                       content_type="multipart/form-data", headers={"Origin": "https://evil.example"})
    assert evil.status_code == 403
    same = client.post("/api/upload/", data={"file": (io.BytesIO(b"a\n1\n"), "x.csv")},
                       content_type="multipart/form-data", headers={"Origin": "http://localhost"})
    assert same.status_code == 201
    curl_like = _upload(client, "x.csv", b"a\n1\n")
    assert curl_like.status_code == 201


def test_listed_origin_is_allowed_to_post(client, monkeypatch):
    from config import Config
    monkeypatch.setattr(Config, "allowed_origins_list", classmethod(lambda cls: ["https://app.example"]))
    r = client.post("/api/upload/", data={"file": (io.BytesIO(b"a\n1\n"), "x.csv")},
                    content_type="multipart/form-data", headers={"Origin": "https://app.example"})
    assert r.status_code == 201


def test_request_paths_cannot_forge_log_lines(client):
    records = []
    class H(logging.Handler):
        def emit(self, record): records.append(record.getMessage())
    h = H(); lg = logging.getLogger("omixa.app"); lg.addHandler(h); lg.setLevel(logging.INFO)
    try:
        client.get("/nope%0AINFO%20forged%20entry")
    finally:
        lg.removeHandler(h)
    assert records and all("\n" not in m for m in records)


def test_error_responses_never_leak_internals(client):
    for path in ("/api/report/zzz", "/api/download/zzz", "/api/does-not-exist"):
        body = client.get(path).get_data(as_text=True)
        assert "Traceback" not in body and "/home/" not in body and "site-packages" not in body


# ============================================================ export safety
def test_formula_payloads_are_defused_in_cells_and_headers():
    from export.exporter import sanitize_formula_injection, _defuse_cell
    df = pd.DataFrame({"=HYPERLINK(\"http://evil\")": ["=1+1", "@SUM(A1)", "+cmd|' /C calc'!A0", "-2+3", "ok"]})
    clean = sanitize_formula_injection(df)
    assert all(not str(v).startswith(("=", "@")) for v in clean.iloc[:, 0])
    assert _defuse_cell("=HYPERLINK(1)").startswith("'")


def test_plain_numbers_and_e164_phones_are_not_corrupted_by_the_defuser():
    """-5 and +2348061234567 were prefixed with an apostrophe, silently corrupting real data."""
    from export.exporter import _defuse_cell
    for ok in ("-5", "+2348061234567", "-0.25", "+1"):
        assert _defuse_cell(ok) == ok
    for bad in ("+1+1", "-5 apples", "=1", "@x", "-cmd|x"):
        assert _defuse_cell(bad).startswith("'")


def test_leading_zero_identifiers_survive_read_clean_export(tmp_path):
    from processing.pipeline import read_source
    p = tmp_path / "ids.csv"
    p.write_text("phone_num,account_no,zip_code\n0803317157,00012345,02134\n0701234567,00099999,02135\n")
    df = read_source(str(p))
    out, _ = apply_rules(df)
    assert out["phone_num"].tolist() == ["0803317157", "0701234567"]
    assert out["account_no"].tolist() == ["00012345", "00099999"]
    assert out["zip_code"].tolist() == ["02134", "02135"]


# ============================================================ frontend XSS
def test_frontend_escapes_every_uploaded_string_it_puts_into_innerhtml():
    """Column names and values come from the uploaded file and were interpolated unescaped
    into innerHTML (stored-XSS-style: a header like <img onerror=...> executed in the page)."""
    js = open(os.path.join(ROOT, "static", "js", "clean.js"), encoding="utf-8").read()
    assert "function escHtml" in js
    for expr in ("f.detail", "f.suggestion", "rec.reason", "f.column", "summary.cleaning_summary.text"):
        for m in re.finditer(r"\$\{\s*" + re.escape(expr) + r"\s*\}", js):
            window = js[max(0, m.start() - 12): m.start()]
            assert "escHtml(" in window or "? escHtml" in window, f"unescaped {expr} near: {window!r}"


def test_malicious_column_names_are_defused_in_the_real_csv_and_xlsx_export(tmp_path):
    """Header cells were written raw: a column named =HYPERLINK(...) became a live formula in .xlsx
    (openpyxl turns any string starting with '=' into one) and a formula in CSV viewers."""
    import tempfile
    from export.exporter import export_dataframe
    import openpyxl
    df = pd.DataFrame({'=HYPERLINK("http://evil","x")': ["ok"], "@SUM(1)": ["fine"], "-5": [1]})
    out_dir = tempfile.mkdtemp()

    csv_path = os.path.join(out_dir, "o.csv")
    export_dataframe(df, csv_path, "csv")
    header = open(csv_path, encoding="utf-8").readline()
    assert not header.startswith(("=", "@")) and ",@" not in header and "'=HYPERLINK" in header

    xlsx_path = os.path.join(out_dir, "o.xlsx")
    export_dataframe(df, xlsx_path, "xlsx")
    ws = openpyxl.load_workbook(xlsx_path).active
    for cell in ws[1]:
        assert cell.data_type != "f", f"header cell {cell.coordinate} is a live formula"
