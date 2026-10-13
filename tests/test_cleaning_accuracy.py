"""Regression tests for the data-cleaning accuracy audit.

Each test pins a defect that was reproduced on a messy workbook (IDs like ID-001 next to 3/6/7, negative ages,
32/13/2025, "NGN 30000", mixed capitalisation, duplicates) before it was fixed, or the safety property that the
fix must not break (nothing is guessed, merged, defaulted or dropped silently).

Files go through the REAL API (upload -> process -> download) as CSV and as XLSX, and the downloaded file is what
the assertions look at, never what a rule "should" have done.
"""
import io

import openpyxl
import pandas as pd
import pytest

from cleaning import casing, detectors
from cleaning.validation import classify_column, detect_type, identifier_shape, validate_dataframe
from export.exporter import _defuse_cell


# --------------------------------------------------------------------------------------------- helpers
def _upload(client, payload: bytes, name: str) -> str:
    r = client.post("/api/upload/", data={"file": (io.BytesIO(payload), name)}, content_type="multipart/form-data")
    assert r.status_code in (200, 201), r.get_json()
    return r.get_json()["job_id"]


def _process(client, job, **body):
    r = client.post(f"/api/process/{job}", json=body)
    assert r.status_code == 200, r.get_json()
    return r.get_json()["summary"]


def run_csv(client, text: str, **body):
    """(summary, downloaded CSV as an all-text frame)"""
    job = _upload(client, text.encode("utf-8"), "messy.csv")
    summary = _process(client, job, **body)
    raw = client.get(f"/api/download/{job}").data.decode("utf-8-sig")
    return summary, pd.read_csv(io.StringIO(raw), dtype=str, keep_default_na=False)


def run_xlsx(client, rows, **body):
    """(summary, downloaded sheet as a list of row tuples, header first)"""
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    job = _upload(client, buf.getvalue(), "messy.xlsx")
    summary = _process(client, job, **body)
    out = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/download/{job}").data))
    return summary, list(out.worksheets[0].iter_rows(values_only=True))


def findings(report, issue=None, column=None):
    return [f for f in report["findings"]
            if (issue is None or f["issue"] == issue) and (column is None or f["column"] == column)]


def col_report(summary, name):
    return next(c for c in summary["integrity_report"]["columns"] if c["column"] == name)


def kinds(df, column):
    """{value: (category, code)} as the validator judges each distinct value of a column."""
    res = classify_column(column, df[column])
    return {v: (res["by_value"][v][0], res["by_value"][v][1]) for v in res["by_value"]}


# The shape of the workbook that exposed the failures (as CSV text; the same rows are used for XLSX below).
MESSY_HEADER = ["Record ID", "Name", "Age", "Signup Date", "Amount", "City", "Status", "Email", "Notes"]
MESSY_ROWS = [
    ["ID-001", "john doe", 25, "2024-01-15", "₦45,000", "lagos", "active", "john@x.com", "ok"],
    ["ID-002", "MARY JANE", -4, "15/02/2024", "NGN 30000", "Lagos", "Active", "mary@x.com", "n/a"],
    ["3", "Peter Obi", 0, "32/13/2025", "12,500", "LAGOS", "ACTIVE", "peter@x.com", "N/A"],
    ["ID-004", "Ada Eze", 41, "03/04/2025", "₦ 1,200.50", "Abuja ", "Inactive", "ada@x.com", "fine"],
    ["6", "Ada Eze", 33, "2025-03-02", "30000", "abuja", "inactive", "ada2@x.com", ""],
    ["7", "Sam", 150, "13/02/2025", "NGN30,000", "Port Harcourt", "Pending", "sam@x.com", "null"],
    ["ID-008", "Sam", 25, "2025-01-01", "", "Port-Harcourt", "pending", "sam2@x.com", "NONE"],
    ["ID-001", "john doe", 25, "2024-01-15", "₦45,000", "lagos", "active", "john@x.com", "ok"],   # exact copy of row 1
]


def messy_csv() -> str:
    def cell(v):
        s = str(v)
        return f'"{s}"' if "," in s else s
    return "\n".join(",".join(cell(v) for v in row) for row in [MESSY_HEADER] + MESSY_ROWS) + "\n"


# ============================================================ A. record identifiers
def test_mixed_identifier_formats_are_reported_and_never_rewritten(client):
    s, out = run_csv(client, messy_csv())
    # nothing was invented: no prefix added to 3/6/7, nothing renumbered
    assert out["Record ID"].tolist() == ["ID-001", "ID-002", "3", "ID-004", "6", "7", "ID-008"]
    hit = findings(s["quality_report_before"], "inconsistent_identifier_format", "Record ID")
    assert len(hit) == 1 and hit[0]["affected_count"] == 3          # the 3 bare numbers differ from the ID-### majority
    assert {f["format"] for f in hit[0]["formats_found"]} == {"A-9", "9"}
    assert "column-specific" in hit[0]["suggestion"]
    # still there afterwards, and the column is not claimed to be "successfully cleaned"
    assert findings(s["quality_report"], "inconsistent_identifier_format", "Record ID")
    assert col_report(s, "Record ID")["status"] != "successfully_cleaned"


def test_identifier_shape_ignores_digit_width_and_prefix_case_but_not_structure():
    assert identifier_shape("ID-001") == identifier_shape("id-1000") == "A-9"
    assert identifier_shape("3") == "9" and identifier_shape("C001") == "A9"


def test_consistent_identifiers_with_growing_width_are_not_flagged(client):
    s, out = run_csv(client, "Customer ID,Name\nCUST-998,a\nCUST-999,b\nCUST-1000,c\nCUST-1001,d\n")
    assert not findings(s["quality_report_before"], "inconsistent_identifier_format")
    assert out["Customer ID"].tolist() == ["CUST-998", "CUST-999", "CUST-1000", "CUST-1001"]


def test_identifiers_keep_leading_zeros_as_text(client):
    s, out = run_csv(client, "Account No,Name\n000123,a\n000456,b\n007890,c\n")
    assert out["Account No"].tolist() == ["000123", "000456", "007890"]
    assert not findings(s["quality_report_before"], "inconsistent_identifier_format")


# ============================================================ B. numeric validation
def test_negative_age_is_invalid_and_preserved_zero_age_is_only_suspicious():
    df = pd.DataFrame({"Age": [25, -4, 0, 41, 150, 33]})
    k = kinds(df, "Age")
    assert k[-4][0] == "invalid" and k[-4][1] == "negative_not_allowed"
    assert k[0] == ("suspicious", "zero_age")                    # not invalid, not missing: a person decides
    assert k[150] == ("suspicious", "implausible_age")           # suspicious differs from definitively invalid
    assert k[25][0] == k[41][0] == k[33][0] == "valid"


def test_age_zero_can_be_declared_valid_by_the_user():
    df = pd.DataFrame({"Age": [25, 0, 41, 33]})
    res = classify_column("Age", df["Age"], {"rules": {"Age": {"min": 0}}})
    assert res["by_value"][0][0] == "valid"


def test_bad_ages_are_kept_exactly_and_never_replaced_by_a_default(client):
    s, out = run_csv(client, messy_csv())
    assert out["Age"].tolist() == ["25", "-4", "0", "41", "33", "150", "25"]
    assert s["integrity_report"]["missing_values"]["filled"] == 0
    assert col_report(s, "Age")["missing_after"] == 0           # 0 and -4 are values, not missing


def test_a_negative_age_is_counted_once_not_in_two_findings(client):
    s, _ = run_csv(client, "Age,Name\n25,a\n-4,b\n41,c\n33,d\n29,e\n")
    ages = [f for f in s["quality_report_before"]["findings"] if f["column"] == "Age"]
    assert [f["issue"] for f in ages] == ["impossible_age"]       # was also reported again as invalid_values
    f = ages[0]
    assert f["affected_count"] == 1 and f["invalid_count"] == 1 and f["suspicious_count"] == 0
    assert "negative (invalid)" in f["detail"]


def test_age_finding_separates_invalid_from_suspicious(client):
    s, _ = run_csv(client, "Age,Name\n25,a\n-4,b\n41,c\n150,d\n29,e\n")
    f = findings(s["quality_report_before"], "impossible_age", "Age")[0]
    assert (f["invalid_count"], f["suspicious_count"]) == (1, 1)
    assert "1 negative (invalid)" in f["detail"] and "1 above 120 (suspicious)" in f["detail"]


def test_legitimate_zero_in_quantity_and_amount_is_neither_flagged_nor_missing(client):
    s, out = run_csv(client, "Quantity,Amount,Name\n0,0,a\n2,1500,b\n5,200,c\n1,0,d\n")
    assert out["Quantity"].tolist() == ["0", "2", "5", "1"] and out["Amount"].tolist() == ["0", "1500", "200", "0"]
    for c in ("Quantity", "Amount"):
        r = col_report(s, c)
        assert r["missing_after"] == 0 and r["needs_review_after"] == 0


# ============================================================ C. dates
def test_valid_dates_are_standardised_even_next_to_ambiguous_and_impossible_ones(client):
    s, out = run_csv(client, messy_csv())
    d = out["Signup Date"].tolist()
    assert d[0] == "2024-01-15" and d[1] == "2024-02-15" and d[4] == "2025-03-02" and d[5] == "2025-02-13"  # 15/02/2024 and 13/02/2025 can only be day-first
    assert d[2] == "32/13/2025"            # impossible: kept byte-for-byte, never turned into another date
    assert d[3] == "03/04/2025"            # ambiguous (3 April or 4 March): not guessed
    # only the two cells that were not already ISO needed rewriting; the other valid ones were already right
    assert s["details"]["date_standardization"]["per_column"]["Signup Date"]["changed"] == 2


def test_impossible_and_ambiguous_dates_are_reported_with_a_reason(client):
    s, _ = run_csv(client, messy_csv())
    col = col_report(s, "Signup Date")
    assert col["counts_after"]["invalid"] == 1 and col["counts_after"]["unresolved"] == 1
    assert col["status"] == "partially_cleaned"
    amb = findings(s["quality_report"], "ambiguous_date_format", "Signup Date")
    assert amb and "not possible to tell day-first from month-first" in amb[0]["detail"]
    res = classify_column("d", pd.Series(["32/13/2025"]))
    assert res["by_value"]["32/13/2025"][:2][0] == "invalid"


def test_a_fully_ambiguous_date_column_is_left_completely_alone(client):
    s, out = run_csv(client, "Name,Signup Date\na,03/04/2024\nb,05/06/2024\nc,01/02/2024\nd,07/08/2024\n")
    assert out["Signup Date"].tolist() == ["03/04/2024", "05/06/2024", "01/02/2024", "07/08/2024"]
    assert not s["details"]["date_standardization"]["per_column"]


def test_two_digit_years_are_not_given_a_century(client):
    s, out = run_csv(client, "Name,Signup Date\na,05/01/24\nb,06/02/68\nc,07/03/99\n")
    assert out["Signup Date"].tolist() == ["05/01/24", "06/02/68", "07/03/99"]


def test_excel_date_cells_and_text_dates_end_up_in_one_format():
    from datetime import datetime
    df = pd.DataFrame({"Join Date": [datetime(2024, 1, 15), "2024-02-15", "15/03/2024", "32/13/2025", "04/05/2024"]})
    from cleaning.rules import apply_rules
    out, _ = apply_rules(df.copy(), rules=["date_standardization"])
    assert out["Join Date"].tolist() == ["2024-01-15", "2024-02-15", "2024-03-15", "32/13/2025", "04/05/2024"]


def test_a_real_time_of_day_still_blocks_rewriting_but_midnight_does_not():
    assert detectors.has_time_component(pd.Series(["2024-01-05 10:30:00", "2024-01-06"]))
    assert not detectors.has_time_component(pd.Series(["2024-01-05 00:00:00", "2024-01-06"]))


# ============================================================ D. numeric and currency formats
@pytest.mark.parametrize("raw,expected", [
    ("₦45,000", 45000), ("NGN 30000", 30000), ("12,500", 12500), ("NGN30,000", 30000), ("₦ 1,200.50", 1200.5),
    ("$1,200", 1200), ("(500)", -500), ("12 500", 12500),
])
def test_supported_numeric_formats_are_read_per_cell(raw, expected):
    status, value, _ = detectors.normalize_numeric_cell(raw)
    assert status == "converted" and value == expected


@pytest.mark.parametrize("raw", ["1.234,56", "1,2,3", "45k", "NGN", "abc"])
def test_ambiguous_or_unreadable_numbers_are_never_guessed(raw):
    status, value, _ = detectors.normalize_numeric_cell(raw)
    assert status in ("ambiguous", "not_numeric") and value is None


def test_amounts_become_numbers_in_the_downloaded_xlsx_and_stay_numbers_in_the_report(client):
    s, rows = run_xlsx(client, [MESSY_HEADER] + MESSY_ROWS)
    amounts = [r[4] for r in rows[1:]]
    assert amounts[:6] == [45000, 30000, 12500, 1200.5, 30000, 30000]
    assert amounts[6] is None                                    # the empty cell stays empty, no default invented
    r = col_report(s, "Amount")
    assert r["detected_type"] in ("currency", "numeric")        # was "phone"
    assert r["counts_after"]["invalid"] == 0 and r["needs_review_after"] == 0


def test_cleaned_amount_values_are_not_mistaken_for_phone_numbers():
    cleaned = pd.Series([45000.0, 30000.0, 12500.0, 1200.5, 30000.0, 1234567.0])
    for name in ("Amount", "Value", "Total"):
        assert detect_type(name, cleaned)["type"] in ("numeric", "currency")
    # a column genuinely named like a phone keeps working, including numbers stored as numbers
    assert detect_type("Phone", pd.Series([8031234567, 8031234568, 8031234569, 8031234570]))["type"] == "phone"
    assert detect_type("Amount", pd.Series(["1500000", "2500000", "3500000", "4500000"]))["type"] != "phone"


def test_a_text_column_with_a_few_numbers_is_not_converted_to_numeric(client):
    s, out = run_csv(client, "Name,Reference\na,alpha\nb,12,500\nc,beta\nd,gamma\n".replace("12,500", '"12,500"'))
    assert out["Reference"].tolist() == ["alpha", "12,500", "beta", "gamma"]   # the one number-like text is untouched


def test_unreadable_amounts_are_preserved_and_flagged_while_valid_ones_are_cleaned(client):
    s, out = run_csv(client, 'Name,Amount\na,"₦2,500"\nb,ERROR\nc,3500\nd,4200\ne,NGN 5000\n')
    assert out["Amount"].tolist() == ["2500", "ERROR", "3500", "4200", "5000"]
    assert col_report(s, "Amount")["needs_review_after"] == 1


def test_an_ambiguous_comma_value_is_preserved_when_the_column_uses_decimal_commas(client):
    # "3,45" shows this column writes decimals with a comma, so "1,234" could be 1.234 or 1234: never guessed
    s, out = run_csv(client, 'Name,Amount\na,"3,45"\nb,"1,234"\nc,12\nd,15\ne,20\nf,30\n')
    assert out["Amount"].tolist()[0] == "3.45" and out["Amount"].tolist()[1] == "1,234"
    assert col_report(s, "Amount")["needs_review_after"] == 1


def test_a_column_that_mixes_two_number_conventions_is_not_guessed_at(client):
    # "₦2,500" says comma = thousands, "1.234,56" says comma = decimal: with contradictory evidence nothing is rewritten
    s, out = run_csv(client, 'Name,Amount\na,"₦2,500"\nb,"1.234,56"\nc,3500\nd,4200\n')
    assert out["Amount"].tolist() == ["₦2,500", "1.234,56", "3500", "4200"]


# ============================================================ E. missing values
def test_missing_markers_are_normalised_once_and_zero_is_never_missing(client):
    s, out = run_csv(client, "Name,Score,Notes\na,0,n/a\nb,45,N/A\nc,0,NULL\nd,82,none\ne,67,ok\nf,0,\n")
    assert out["Notes"].tolist() == ["", "", "", "", "ok", ""]
    assert out["Score"].tolist() == ["0", "45", "0", "82", "67", "0"]
    ir = s["integrity_report"]["missing_values"]
    assert ir["converted_from_placeholders"] == 4 and ir["filled"] == 0
    assert col_report(s, "Score")["missing_after"] == 0


def test_cleaning_does_not_introduce_unintended_missing_values(client):
    s, out = run_csv(client, messy_csv())
    before_missing = {c: int(m["originally_missing"]) for c, m in s["missing_tracking"].items()}
    for col, m in s["missing_tracking"].items():
        # every extra empty cell is explained by a converted placeholder, never by a transformation
        assert m["missing_after"] == m["originally_missing"] + m["converted_to_missing"] - m["filled"] - _removed_missing(m)
    assert before_missing["Amount"] == 1 and (out["Amount"] == "").sum() == 1


def _removed_missing(m):
    return 0


# ============================================================ F. categorical values
def test_case_variants_in_a_small_column_are_unified_without_inventing_values(client):
    s, out = run_csv(client, messy_csv())
    assert set(out["Status"]) == {"Active", "Inactive", "Pending"}
    assert set(out["City"]) == {"Lagos", "Abuja", "Port Harcourt"}          # 'Abuja ' trimmed, hyphen/space variant merged
    assert not findings(s["quality_report"], "inconsistent_categories")
    assert not findings(s["quality_report"], "inconsistent_casing", "Status")


def test_a_column_is_not_skipped_just_because_its_spellings_are_many():
    # 7 spellings of 3 labels in 8 rows: counting spellings (7 > 4) wrongly made it look like free text
    df = pd.DataFrame({"Status": ["active", "Active", "ACTIVE", "Inactive", "inactive", "Pending", "pending", "Active"]})
    assert casing.distinct_labels(df["Status"].unique()) == 3
    from cleaning.rules import apply_rules
    out, _ = apply_rules(df.copy(), rules=["case_standardization", "categorical_standardization"])
    assert set(out["Status"]) == {"Active", "Inactive", "Pending"}


@pytest.mark.parametrize("a,b", [("AB+", "AB-"), ("Grade A+", "Grade A-"), ("R&D", "RD"), ("Type 1", "Type 2")])
def test_distinct_categories_are_never_treated_as_one(a, b):
    assert casing.variant_key(a) != casing.variant_key(b)


@pytest.mark.parametrize("a,b", [("Port-Harcourt", "Port Harcourt"), ("I.T.", "IT"), ("Web site", "Website"), ("Male", "MALE")])
def test_genuine_spelling_variants_still_share_a_key(a, b):
    assert casing.variant_key(a) == casing.variant_key(b)


def test_blood_groups_and_grades_are_not_merged_by_the_cleaning_rules(client):
    s, out = run_csv(client, "Patient,Blood Group,Grade\n1,AB+,Grade A+\n2,AB-,Grade A-\n3,AB+,Grade A+\n4,AB-,Grade A-\n"
                             "5,O+,Grade B\n6,O-,Grade B\n7,AB+,Grade A+\n8,AB-,Grade B\n")
    assert out["Blood Group"].tolist() == ["AB+", "AB-", "AB+", "AB-", "O+", "O-", "AB+", "AB-"]
    assert out["Grade"].tolist() == ["Grade A+", "Grade A-", "Grade A+", "Grade A-", "Grade B", "Grade B", "Grade A+", "Grade B"]
    assert not findings(s["quality_report_before"], "inconsistent_categories")


def test_a_probable_typo_is_flagged_not_merged(client):
    s, out = run_csv(client, "Name,City\n" + "\n".join(f"p{i},Lagos" for i in range(6)) + "\nq,Lagoss\n")
    assert "Lagoss" in set(out["City"]) and out["City"].tolist().count("Lagos") == 6
    assert findings(s["quality_report"], "possible_misspelling", "City")


# ============================================================ G. duplicates
def test_only_exact_duplicate_rows_are_removed_and_first_is_kept(client):
    s, out = run_csv(client, messy_csv())
    assert len(out) == 7 and out["Record ID"].tolist().count("ID-001") == 1
    assert s["integrity_report"]["duplicates"]["exact_rows_removed"] == 1
    assert s["change_log"]["removed_rows_total"] == 1


def test_same_name_with_different_email_is_not_a_duplicate(client):
    s, out = run_csv(client, "Name,Email,City\nAda Eze,ada@x.com,Lagos\nAda Eze,ada2@x.com,Lagos\nBob,bob@x.com,Abuja\n")
    assert len(out) == 3
    ir = s["integrity_report"]["duplicates"]
    assert ir["exact_rows_removed"] == 0 and ir["potential"]["groups"] == 0 and ir["by_key"]["groups"] == 0


def test_same_email_with_different_details_is_flagged_for_review_never_removed(client):
    s, out = run_csv(client, "Name,Email,City\nSam,sam@x.com,Lagos\nSamuel,sam@x.com,Abuja\nBob,bob@x.com,Abuja\n")
    assert len(out) == 3
    assert s["integrity_report"]["duplicates"]["potential"]["groups"] == 1


# ============================================================ H. quality reports
def test_before_and_after_reports_describe_the_actual_data(client):
    s, out = run_csv(client, messy_csv())
    before, after = s["quality_report_before"], s["quality_report"]
    assert before["row_count"] == 8 and after["row_count"] == 7
    # things the fixers resolved are gone, things that need a person are still reported
    for issue, col in (("inconsistent_categories", "Status"), ("inconsistent_casing", "City"), ("duplicate_rows", None)):
        assert findings(before, issue, col) and not findings(after, issue, col)
    for issue, col in (("inconsistent_identifier_format", "Record ID"), ("ambiguous_date_format", "Signup Date"),
                       ("impossible_age", "Age")):
        assert findings(before, issue, col) and findings(after, issue, col)
    assert after["counts"]["warning"] < before["counts"]["warning"]
    cmp = s["quality_comparison"]["overall"]
    assert cmp["after_without_cap"] >= cmp["before_without_cap"]            # the uncapped score is not hidden by the cap


def test_the_final_report_matches_an_independent_validation_of_the_downloaded_file(client):
    s, out = run_csv(client, messy_csv())
    independent = validate_dataframe(out.replace("", None).convert_dtypes(), None, include_examples=False)
    # compare per column on the columns whose validation does not depend on dtype inference
    for name in ("Record ID", "Signup Date", "City", "Status"):
        ours = col_report(s, name)["counts_after"]
        theirs = next(c for c in independent["columns"] if c["column"] == name)["counts"]
        assert {k: ours[k] for k in ("invalid", "suspicious", "unresolved")} == \
               {k: theirs[k] for k in ("invalid", "suspicious", "unresolved")}, name


def test_issue_summary_separates_corrected_remaining_and_review(client):
    s, _ = run_csv(client, messy_csv())
    sm = s["integrity_report"]["issue_summary"]
    assert sm["detected"] > sm["remaining"] > 0
    assert sm["automatically_corrected"] == sm["detected"] - sm["remaining"] + sm["introduced_by_cleaning"] - sm["removed_with_duplicate_rows"]
    assert sm["requires_review"] + sm["cannot_be_corrected_automatically"] == sm["remaining"]
    assert sm["cannot_be_corrected_automatically"] == 2                      # -4 and 32/13/2025
    assert sm["cells_normalised"] > 0 and set(sm["definitions"]) >= {"detected", "remaining", "requires_review"}
    assert s["integrity_report"]["status"] == "partially_cleaned" and s["integrity_report"]["fully_clean"] is False


def test_issues_that_vanish_only_because_a_duplicate_row_was_removed_are_not_counted_as_fixed(client):
    # the same bad row twice: one copy is removed, the other keeps its problem. Nothing was corrected.
    s, out = run_csv(client, "Name,Age\na,25\nb,-4\nb,-4\nc,41\nd,33\n")
    assert len(out) == 4 and out["Age"].tolist() == ["25", "-4", "41", "33"]
    iv = s["integrity_report"]["invalid_values"]
    assert iv["detected"] == 2 and iv["requiring_manual_review"] == 1
    assert iv["resolved"] == 0 and iv["removed_with_duplicate_rows"] == 1
    assert s["integrity_report"]["issue_summary"]["automatically_corrected"] == 0


def test_a_fix_in_one_column_is_not_netted_against_a_new_problem_in_another(client):
    s, _ = run_csv(client, messy_csv())
    cols = s["integrity_report"]["columns"]
    assert sum(c.get("invalid_resolved", 0) for c in cols) == s["integrity_report"]["invalid_values"]["resolved"]
    assert all(c.get("invalid_introduced", 0) == 0 for c in cols)


def test_a_clean_file_is_reported_clean_and_unchanged(client):
    s, out = run_csv(client, "Customer ID,Name,Age\nC-001,Ada Eze,25\nC-002,Bob Obi,31\nC-003,Cy Eke,44\n")
    assert out.values.tolist() == [["C-001", "Ada Eze", "25"], ["C-002", "Bob Obi", "31"], ["C-003", "Cy Eke", "44"]]
    ir = s["integrity_report"]
    assert ir["status"] == "successfully_cleaned" and ir["fully_clean"] is True
    assert ir["issue_summary"]["detected"] == 0 and ir["issue_summary"]["cells_normalised"] == 0


# ============================================================ I. CSV and XLSX parity, and preservation
def test_csv_and_xlsx_give_the_same_cleaned_values_and_the_same_open_issues(client):
    s_csv, out_csv = run_csv(client, messy_csv())
    s_xlsx, rows = run_xlsx(client, [MESSY_HEADER] + MESSY_ROWS)
    assert len(rows) == 8 and s_xlsx["rows_out"] == s_csv["rows_out"] == 7
    for j, name in enumerate(MESSY_HEADER):
        if name == "Amount":
            continue                                    # numbers in XLSX, text in CSV: compared below
        xlsx_col = ["" if r[j] is None else str(r[j]) for r in rows[1:]]
        assert xlsx_col == out_csv[name].tolist(), name
    assert [float(x) for x in out_csv["Amount"] if x] == [float(r[4]) for r in rows[1:] if r[4] is not None]
    a, b = s_csv["integrity_report"]["issue_summary"], s_xlsx["integrity_report"]["issue_summary"]
    assert {k: a[k] for k in ("detected", "remaining", "requires_review", "cannot_be_corrected_automatically")} == \
           {k: b[k] for k in ("detected", "remaining", "requires_review", "cannot_be_corrected_automatically")}


def test_xlsx_keeps_identifiers_dates_and_unsafe_values_exactly_where_correction_is_unsafe(client):
    s, rows = run_xlsx(client, [MESSY_HEADER] + MESSY_ROWS)
    body = {r[0]: r for r in rows[1:]}
    assert body["3"][2] == 0 and body["ID-002"][2] == -4 and body["7"][2] == 150          # ages untouched
    assert body["3"][3] == "32/13/2025" and body["ID-004"][3] == "03/04/2025"             # dates untouched
    assert s["export_verification"]["verified"] is True


def test_every_row_is_accounted_for_nothing_is_dropped_silently(client):
    s, out = run_csv(client, messy_csv())
    assert s["rows_in"] - s["rows_out"] == s["integrity_report"]["rows"]["removed"] == s["change_log"]["removed_rows_total"] == 1


# ============================================================ export guard
def test_a_lone_dash_is_not_prefixed_but_real_formulas_still_are():
    for ok in ("-", "+", "-5", "+2348061234567"):
        assert _defuse_cell(ok) == ok
    for bad in ("=1+1", "@SUM(A1)", "-cmd|x", "+1+1", "--1"):
        assert _defuse_cell(bad).startswith("'")


def test_dash_placeholder_survives_the_export_unchanged(client):
    s, out = run_csv(client, "Name,Notes\na,-\nb,ok\nc,fine\n")
    assert out["Notes"].tolist() == ["-", "ok", "fine"]
