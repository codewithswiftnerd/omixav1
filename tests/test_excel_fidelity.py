"""What Omixa does and does NOT preserve when it cleans an .xlsx (see docs/EXCEL_FIDELITY.md).
Preserved: sheet names/other sheets, workbook properties, per-column font + width, bold header,
hidden columns, a frozen header row. NOT preserved on the cleaned sheet: formulas (values only),
merged cells, hyperlinks, comments, conditional formatting, data validation, tables, row heights,
cell fills, number formats. These tests pin both halves so the UI/FAQ can stay honest."""
import io
import os
import zipfile

import openpyxl
import pytest
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table

from config import Config
from processing.pipeline import run_pipeline
from tests.conftest import make_csv

JOB = "22222222-2222-4222-8222-222222222222"


def _workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Amount", "Site", "Secret"])
    for i in range(2, 6):
        ws.append([f"  Ada{i} ", i * 10, None, "x"])
    ws["C2"] = "site"
    ws["C2"].hyperlink = "https://example.com"
    ws["A1"].font = Font(bold=True)
    ws["A2"].fill = PatternFill("solid", fgColor="FFFF00")
    ws.column_dimensions["A"].width = 33
    ws.column_dimensions["D"].hidden = True
    ws.freeze_panes = "A2"
    ws["B3"].comment = Comment("note", "me")
    ws.merge_cells("A8:B8")
    ws["A8"] = "merged"
    ws.add_table(Table(displayName="T1", ref="A1:D5"))
    ws["E2"] = "=B2*2"
    ws["E1"] = "Calc"
    wb.create_sheet("Other")["A1"] = "keep me"
    wb.properties.title = "My Title"
    wb.save(path)


@pytest.fixture
def cleaned(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "TEMP_DIR", str(tmp_path))
    os.makedirs(tmp_path / JOB)
    _workbook(str(tmp_path / JOB / "source.xlsx"))
    run_pipeline(JOB, rules=None)
    return openpyxl.load_workbook(str(tmp_path / JOB / "cleaned.xlsx"))


def test_preserved_features(cleaned):
    ws = cleaned["Data"]
    assert cleaned.sheetnames == ["Data", "Other"] and cleaned["Other"]["A1"].value == "keep me"
    assert cleaned.properties.title == "My Title"
    assert ws["A1"].font.b is True
    assert ws.column_dimensions["A"].width == 33
    assert ws.column_dimensions["D"].hidden is True
    assert ws.freeze_panes == "A2"
    assert ws["A2"].value == "Ada2"  # whitespace cleaned, value kept


def test_documented_losses(cleaned):
    """Known, documented limitations. If one of these starts passing the other way, update the docs/UI copy."""
    ws = cleaned["Data"]
    assert ws.merged_cells.ranges == set() or "A8:B8" not in {str(r) for r in ws.merged_cells.ranges}
    assert ws["C2"].hyperlink is None
    assert ws["B3"].comment is None
    assert not ws.tables
    assert not any(isinstance(c.value, str) and c.value.startswith("=") and c.value != "=" for r in ws.iter_rows() for c in r
                   if c.value and "B2" in str(c.value))  # formulas are not carried over (values only)


def test_formula_injection_is_defused(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "TEMP_DIR", str(tmp_path))
    os.makedirs(tmp_path / JOB)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Note"])
    ws.append(["Ada", "=HYPERLINK(\"http://evil\",\"x\")"])
    ws.append(["Bob", "+1+1"])
    wb.save(str(tmp_path / JOB / "source.xlsx"))
    run_pipeline(JOB, rules=None)
    out = openpyxl.load_workbook(str(tmp_path / JOB / "cleaned.xlsx"))["Sheet"]
    for r in (2, 3):
        v = out.cell(row=r, column=2).value
        assert not str(v).startswith(("=", "+", "-", "@"))


def test_macro_workbooks_are_rejected(client):
    r = client.post("/api/upload/", data={"file": (io.BytesIO(b"PK\x03\x04"), "m.xlsm")}, content_type="multipart/form-data")
    assert r.status_code == 400
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<x/>")
        z.writestr("xl/vbaProject.bin", b"macro")
    r = client.post("/api/upload/", data={"file": (io.BytesIO(buf.getvalue()), "sneaky.xlsx")}, content_type="multipart/form-data")
    assert r.status_code == 400
